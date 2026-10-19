#!/usr/bin/env python3
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0F, 0x2A, 0x3D)
TERR = RGBColor(0x0B, 0x6E, 0x4F)
PAPER = RGBColor(0xEE, 0xF3, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x12, 0x20, 0x2B)
MUTED = RGBColor(0x4D, 0x5C, 0x58)
OK = RGBColor(0x1F, 0x7A, 0x4D)
BAD = RGBColor(0xA6, 0x1F, 0x2B)
CODE = RGBColor(0x10, 0x20, 0x29)
CREAM = RGBColor(0xE7, 0xF3, 0xEA)

W = Inches(13.333)
H = Inches(7.5)
HERE = Path(__file__).resolve().parent
SHOTS = HERE / "gorseller"
OUT = HERE / "Kampus_Kirtasiye_Yazilim_Testi_Sunumu.pptx"


def run(r, size=18, bold=False, color=INK, font="Calibri"):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font


def box(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.line.fill.background()
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    return sh


def roundb(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.adjustments[0] = 0.08
    sh.line.fill.background()
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    return sh


def text(slide, l, t, w, h, s, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    run(r, size, bold, color)
    return tb


def footer(slide, no):
    box(slide, 0, Inches(7.22), W, Inches(0.28), NAVY)
    text(slide, Inches(0.4), Inches(7.22), Inches(10.5), Inches(0.28),
         "Kampüs Kırtasiye  ·  İTÜ BİDB + TechTarget  ·  Rıdvan Çağlar Mutlu", 11, False, CREAM)
    text(slide, Inches(11.6), Inches(7.22), Inches(1.4), Inches(0.28), f"{no:02d}", 11, True, TERR, PP_ALIGN.RIGHT)


def chrome(slide, no, title):
    box(slide, 0, 0, W, H, PAPER)
    box(slide, 0, 0, W, Inches(0.12), TERR)
    box(slide, 0, Inches(0.12), W, Inches(0.92), NAVY)
    text(slide, Inches(0.45), Inches(0.28), Inches(12), Inches(0.6), title, 24, True, WHITE)
    footer(slide, no)


def bullets(slide, l, t, w, items, size=18):
    tb = slide.shapes.add_textbox(l, t, w, Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        r = p.add_run()
        r.text = "▸  " + item
        run(r, size, False, INK)


def table_slide(slide, rows, l=Inches(0.4), t=Inches(1.25), w=Inches(12.5), h=Inches(4.6)):
    tbl = slide.shapes.add_table(len(rows), len(rows[0]), l, t, w, h).table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run_ = p.add_run()
            run_.text = val
            header = r == 0
            run(run_, 12, header, WHITE if header else INK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if header else (WHITE if r % 2 else RGBColor(0xDC, 0xE8, 0xE3))
    return tbl


def code_box(slide, l, t, w, h, s):
    roundb(slide, l, t, w, h, CODE)
    tb = slide.shapes.add_textbox(l + Inches(0.2), t + Inches(0.14), w - Inches(0.35), h - Inches(0.24))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = s
    run(r, 13, False, CREAM, "Consolas")


def pic(slide, name, l, t, w, h):
    p = SHOTS / name
    if p.exists():
        slide.shapes.add_picture(str(p), l, t, w, h)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    box(s, 0, 0, W, H, NAVY)
    box(s, 0, 0, Inches(0.22), H, TERR)
    text(s, Inches(0.7), Inches(1.15), Inches(12), Inches(0.4), "YAZILIM TESTİ DERSİ  ·  İKİ PARÇA TESLİM", 15, True, TERR)
    text(s, Inches(0.7), Inches(1.65), Inches(12), Inches(1.3), "Önce web sitesi.\nSonra onu test eden uygulama.", 40, True, WHITE)
    text(s, Inches(0.7), Inches(4.2), Inches(11.5), Inches(1.8),
         "Kampüs Kırtasiye  +  Test uygulaması (Playwright)\n"
         "Kuram: İTÜ BİDB Yazılım Testi ve Test Süreçleri  ·  TechTarget Software testing\n"
         "Rıdvan Çağlar Mutlu  ·  2514216055  ·  18 Ekim 2026", 18, False, CREAM)
    text(s, Inches(12.2), Inches(6.7), Inches(0.8), Inches(0.35), "01", 16, True, TERR, PP_ALIGN.RIGHT)

    s = prs.slides.add_slide(blank)
    chrome(s, 2, "Teslim neyden oluşur?")
    cards = [
        ("1. Web sitesi", "Kampüs Kırtasiye: giriş, 6 ürün, sepet, 10 haneli telefon, demo sipariş. V2 teslim, V1 hata kanıtı."),
        ("2. Test uygulaması", "Hazırlık, 15 senaryo, V1/V2 sonuç, Playwright otomasyonu. Siteyi tarayıcıdan dener."),
        ("3. Bu sunum", "İTÜ sürecini (hazırlık → dinamik test → sonlandırma) kendi sitemize bağlar."),
    ]
    for i, (a, b) in enumerate(cards):
        y = Inches(1.3 + i * 1.8)
        roundb(s, Inches(0.5), y, Inches(12.3), Inches(1.6), WHITE)
        text(s, Inches(0.75), y + Inches(0.2), Inches(11.8), Inches(0.45), a, 22, True, NAVY)
        text(s, Inches(0.75), y + Inches(0.7), Inches(11.8), Inches(0.7), b, 16, False, MUTED)

    s = prs.slides.add_slide(blank)
    chrome(s, 3, "İTÜ BİDB: Test nedir?")
    bullets(s, Inches(0.5), Inches(1.3), Inches(12.2), [
        "Test: sistemi elle veya otomatik deneyerek gereksinimi doğrulamak, ya da beklenen ile gözlenen farkı bulmak.",
        "Yazılım testi: sonsuz çalışma alanından sınırlı ve uygun senaryolar seçerek dinamik doğrulama yapmak.",
        "Kaliteli yazılım: kabul edilebilir hatasız, bütçe ve zamanda biten, gereksinimi karşılayan, sürdürülebilir.",
        "Bu derste beklenen: giriş kabul/ret, 4 ürün filtresi, 398,80 TL, 10 haneli telefon, boş sepet yasağı.",
    ], 18)

    s = prs.slides.add_slide(blank)
    chrome(s, 4, "İTÜ: Test neden yapılır?")
    bullets(s, Inches(0.5), Inches(1.3), Inches(12.2), [
        "Yazılımdaki hatalar kaçınılmazdır; sonuç para, zaman veya güven kaybı olabilir.",
        "Müşteriye (burada öğretmene) vermeden kaliteden emin olmak.",
        "Hatayı erken yakalamak: V1’de T04 ve T12 bulundu, V2’de kapatıldı.",
        "Yeniden iş ve düzeltme maliyetini düşürmek.",
        "TechTarget aynı gerekçeyi Starbucks kasa çökmesi ve Nissan hava yastığı örneğiyle anlatır.",
    ], 18)

    s = prs.slides.add_slide(blank)
    chrome(s, 5, "İTÜ test süreci: üç aşama")
    cols = [
        ("Hazırlık", "Doküman, kapsam, risk, ortam, görev, bitirme ölçütü, test case, test planı."),
        ("Dinamik test", "Kod çalışırken birim, tümleyim, regresyon, performans, kabul, kara/beyaz kutu."),
        ("Sonlandırma", "Ölçüt sağlandı mı? Hatalar kapandı mı? Kabul testine açılır mı?"),
    ]
    for i, (a, b) in enumerate(cols):
        x = Inches(0.4 + i * 4.25)
        roundb(s, x, Inches(1.4), Inches(4.05), Inches(4.6), WHITE)
        text(s, x + Inches(0.25), Inches(1.65), Inches(3.55), Inches(0.8), f"{i+1}. {a}", 22, True, TERR)
        text(s, x + Inches(0.25), Inches(2.6), Inches(3.55), Inches(3.0), b, 16, False, INK)

    s = prs.slides.add_slide(blank)
    chrome(s, 6, "TechTarget: Yazılım testi nedir, neden önemli?")
    bullets(s, Inches(0.5), Inches(1.3), Inches(12.2), [
        "İşlevi değerlendirmek: hata, boşluk ve beklenen sonuç uyuşmazlığını aramak.",
        "Canlıya çıkmadan hata düzeltmek: kalite, güven ve maliyet.",
        "Erken hata = ucuz hata. Biz T04/T12’yi teslimden önce yakaladık.",
        "Otomasyon: aynı senaryoyu tekrar tekrar koşmak (Playwright).",
        "Regresyon atlanmamalı: düzeltme eski 13 senaryoyu bozmamalı.",
    ], 18)

    s = prs.slides.add_slide(blank)
    chrome(s, 7, "Test türleri: kaynakta var, biz hangisini yaptık?")
    table_slide(s, [
        ["Tür (kaynak)", "Bu projede", "Kanıt"],
        ["Kara kutu / fonksiyonel (İTÜ + TechTarget)", "Evet, asıl yöntem", "15 Playwright senaryosu"],
        ["Beyaz kutu (İTÜ)", "Kısmen", "Bozuk option ve PHP if incelendi"],
        ["Tümleyim (İTÜ / TechTarget)", "Evet", "Giriş + sepet + form birlikte"],
        ["Regresyon", "Evet", "Aynı 15 test V2’de yeniden"],
        ["Kabul (UAT)", "Ders ölçütü", "15/15 ve canlı gösterim"],
        ["Birim / performans / güvenlik", "Hayır", "Kapsam dışı, slayt 22"],
    ], h=Inches(5.2))

    s = prs.slides.add_slide(blank)
    chrome(s, 8, "Hazırlık: ne test edilir, ne edilmez?")
    bullets(s, Inches(0.5), Inches(1.25), Inches(6.0), [
        "Edilecek: giriş, kilitli hesap, 6 ürün, kırtasiye filtresi.",
        "Edilecek: sepet adedi, 14,90 TL kargo, 398,80 TL.",
        "Edilecek: 9 / 10 / 11 haneli telefon.",
        "Edilecek: boş sepet yasağı, çıkış sonrası koruma.",
    ], 17)
    bullets(s, Inches(6.7), Inches(1.25), Inches(6.0), [
        "Edilmeyecek: yük / performans testi.",
        "Edilmeyecek: sızma / XSS taraması.",
        "Edilmeyecek: mobil ve gerçek ödeme.",
        "Bitirme: V2’de 15/15, T04 ve T12 kapanmış.",
    ], 17)

    s = prs.slides.add_slide(blank)
    chrome(s, 9, "Parça 1 — Web sitesi: Kampüs Kırtasiye")
    pic(s, "teslim.png", Inches(0.4), Inches(1.2), Inches(7.3), Inches(4.8))
    bullets(s, Inches(7.9), Inches(1.3), Inches(5.0), [
        "Kapı sayfası iki parçayı ayırır.",
        "Site: /v2/  (teslim)",
        "Hatalı kanıt: /v1/",
        "Giriş: ridvan / Defter2026",
    ], 16)

    s = prs.slides.add_slide(blank)
    chrome(s, 10, "Sınanan uygulama: öğrenci rafı")
    pic(s, "raf.png", Inches(0.35), Inches(1.15), Inches(12.6), Inches(5.8))

    s = prs.slides.add_slide(blank)
    chrome(s, 11, "Parça 2 — Test uygulaması")
    pic(s, "test-surec.png", Inches(0.35), Inches(1.15), Inches(12.6), Inches(5.8))

    s = prs.slides.add_slide(blank)
    chrome(s, 12, "Test case: girdi verilince çıktı ne olmalı?")
    table_slide(s, [
        ["ID", "Girdi", "Beklenen çıktı"],
        ["T02", "ridvan / yanlis123", "Şifre hatalı, sayfa=giris"],
        ["T04", "Kategori = kirtasiye", "4 ürün listelenir"],
        ["T10", "Defter + USB", "398,80 TL (383,90 + 14,90)"],
        ["T11", "0555123456", "Sipariş tamamlanır"],
        ["T12", "0 ürün + Teslimata geç", "Sepette kal, uyarı göster"],
        ["T13 / T14", "9 veya 11 hane", "Telefon 10 rakam olmalıdır"],
    ], h=Inches(5.0))

    s = prs.slides.add_slide(blank)
    chrome(s, 13, "Dinamik test: Playwright siteye nasıl bağlanır?")
    code_box(s, Inches(0.45), Inches(1.25), Inches(12.4), Inches(4.6),
             "use: {\n"
             "  baseURL: 'http://localhost/kampus-kirtasiye/v2/',\n"
             "  viewport: { width: 1366, height: 900 },\n"
             "  screenshot: 'only-on-failure',\n"
             "  trace: 'retain-on-failure'\n"
             "}\n"
             "test.beforeEach(async ({ page }) => {\n"
             "  await page.goto('index.php?sayfa=giris');\n"
             "});")
    text(s, Inches(0.5), Inches(6.05), Inches(12), Inches(0.7),
         "Aynı 15 test V1 ve V2’ye, 1 worker ve 0 retry ile koşuldu.", 16, False, MUTED)

    s = prs.slides.add_slide(blank)
    chrome(s, 14, "Kara kutu örneği: tıkla, sonra expect et")
    code_box(s, Inches(0.45), Inches(1.25), Inches(12.4), Inches(4.3),
             "test('T02 Yanlis sifre reddedilir', async ({ page }) => {\n"
             "  await giris(page, 'ridvan', 'yanlis123');\n"
             "  await expect(page.getByTestId('error'))\n"
             "    .toHaveText('Kullanıcı adı veya şifre hatalı.');\n"
             "  await expect(page).toHaveURL(/sayfa=giris/);\n"
             "});")
    text(s, Inches(0.5), Inches(5.75), Inches(12), Inches(0.9),
         "İşlem tıklamaktır. Doğrulama expect’tir. Kodun içini okumadan dışarıdan sınarız.", 16, False, INK)

    s = prs.slides.add_slide(blank)
    chrome(s, 15, "V1 dinamik test: beklenen ≠ gözlenen")
    roundb(s, Inches(0.45), Inches(1.4), Inches(3.9), Inches(2.4), WHITE)
    text(s, Inches(0.7), Inches(1.6), Inches(3.5), Inches(0.9), "13 geçti / 2 kaldı", 22, True, BAD)
    text(s, Inches(0.7), Inches(2.5), Inches(3.5), Inches(1.0), "T04 ve T12.\nSüre 33,9 sn.", 15, False, MUTED)
    bullets(s, Inches(4.7), Inches(1.4), Inches(8), [
        "T04: kirtasiye seçeneği DOM’da yok (bozuk option).",
        "T12: boş sepet teslimat sayfasını açtı.",
        "İTÜ tanımı burada somut: fark bulundu, test durmadı.",
        "Kanıt: hata metni, ekran görüntüsü, trace.",
    ], 17)

    s = prs.slides.add_slide(blank)
    chrome(s, 16, "H01: HTML option (beyaz kutu inceleme)")
    code_box(s, Inches(0.45), Inches(1.25), Inches(12.4), Inches(4.4),
             "<!-- V1 -->\n"
             "</option value=\"kirtasiye\">Yalnız kırtasiye\n\n"
             "<!-- V2 -->\n"
             "<option value=\"kirtasiye\">Yalnız kırtasiye</option>\n\n"
             "await page.getByLabel('Kategori').selectOption('kirtasiye');\n"
             "await expect(page.getByTestId('product')).toHaveCount(4);")
    text(s, Inches(0.5), Inches(5.9), Inches(12), Inches(0.8),
         "Kara kutu T04 düştü. Nedeni için HTML’e bakıldı. Düzeltme site kodunda.", 16, False, INK)

    s = prs.slides.add_slide(blank)
    chrome(s, 17, "H02: Boş sepet — beklenen uyarı, gözlenen form")
    pic(s, "v1-bos-teslimat.png", Inches(0.35), Inches(1.2), Inches(6.3), Inches(4.15))
    pic(s, "sepet-bos-v2.png", Inches(6.75), Inches(1.2), Inches(6.2), Inches(4.15))
    text(s, Inches(0.4), Inches(5.5), Inches(6.2), Inches(0.5), "V1 gözlenen: teslimat açıldı", 16, True, BAD)
    text(s, Inches(6.75), Inches(5.5), Inches(6.2), Inches(0.5), "V2 gözlenen: sepet + uyarı", 16, True, OK)

    s = prs.slides.add_slide(blank)
    chrome(s, 18, "Düzeltme sunucuda: kontrolü düğmeye bırakma")
    code_box(s, Inches(0.45), Inches(1.25), Inches(12.4), Inches(4.4),
             "if (in_array($sayfa, ['teslimat', 'onay'], true)\n"
             "    && sepet_adet() === 0) {\n"
             "    flash('Sepetiniz boş. Teslimata geçmeden önce bir ürün ekleyin.');\n"
             "    git('sepet');\n"
             "}\n"
             "// teslimata_gec ve siparis POST’una da aynı kural eklendi.")
    text(s, Inches(0.5), Inches(5.9), Inches(12), Inches(0.8),
         "TechTarget: hatayı teslimden önce kapat. İTÜ: dinamik test, hatalar çözülmeden bitmez.", 16, False, INK)

    s = prs.slides.add_slide(blank)
    chrome(s, 19, "Regresyon: aynı 15 test, düzeltilmiş site")
    table_slide(s, [
        ["Koşu", "Geçti", "Kaldı", "Süre"],
        ["V1 dinamik test", "13", "2", "33,9 sn"],
        ["V2 regresyon", "15", "0", "7,8 sn"],
        ["Sonlandırma ölçütü", "15/15", "sağlandı", "V2 teslim"],
    ], h=Inches(3.1))
    text(s, Inches(0.5), Inches(4.7), Inches(12.3), Inches(1.8),
         "Süre farkı performans testi değildir. V1, T04’te olmayan seçeneği bekler.\n"
         "Regresyonun anlamı: T04 ve T12 kapandı, diğer 13 senaryo bozulmadı.", 16, False, MUTED)

    s = prs.slides.add_slide(blank)
    chrome(s, 20, "Test uygulamasındaki sonuç ekranı")
    pic(s, "test-sonuc.png", Inches(0.35), Inches(1.15), Inches(12.6), Inches(5.8))

    s = prs.slides.add_slide(blank)
    chrome(s, 21, "Derste canlı sıra")
    steps = [
        "1. http://localhost/kampus-kirtasiye/  → iki parçayı göster",
        "2. Web sitesi V2: giriş, defter ekle, 398,80 TL, 10 haneli telefon",
        "3. Test uygulaması: süreç, senaryo tablosu, V1/V2 sonuç",
        "4. PLAYWRIGHT_AC.cmd: V1’de T04/T12, V2’de 15/15",
    ]
    for i, st in enumerate(steps):
        roundb(s, Inches(0.5), Inches(1.35 + i * 1.3), Inches(12.3), Inches(1.1), WHITE)
        text(s, Inches(0.75), Inches(1.55 + i * 1.3), Inches(11.8), Inches(0.7), st, 20, False, INK)

    s = prs.slides.add_slide(blank)
    chrome(s, 22, "Sonlandırma, sınır ve kaynaklar")
    bullets(s, Inches(0.5), Inches(1.3), Inches(12.2), [
        "Ölçüt sağlandı: V2 15/15, T04 ve T12 kapandı, regresyon yeşil.",
        "15/15 yalnızca seçilen kapsamı gösterir; ürünün ‘hatasız’ olduğu anlamına gelmez.",
        "Yük, güvenlik ve mobil test yapılmadı (İTÜ / TechTarget’de var, biz seçmedik).",
        "İTÜ BİDB: Yazılım Testi ve Test Süreçleri (2013).",
        "TechTarget: Software testing — definition, types and importance.",
    ], 18)

    prs.save(OUT)
    root = Path("/workspace") / OUT.name
    root.write_bytes(OUT.read_bytes())
    print("saved", OUT, OUT.stat().st_size)


if __name__ == "__main__":
    build()
